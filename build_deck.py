"""Builds the AQUA SIGHT pitch deck as a .pptx (same layout as aqua-sight-pitch-deck.gs).

Design canvas is 1333.33 x 750 units on a 13.333in x 7.5in slide (1 unit = 0.01in).
"""
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

U = 9144  # EMU per design unit
FS = 0.72  # design font size -> points on a 960pt-wide slide

NAVY, TEAL, AQUA = "0B3D5C", "1F7A8C", "3FC1C9"
BG, SLATE, WHITE = "F7FAFC", "4A5568", "FFFFFF"
GREEN, YELLOW, RED = "2F9E44", "F2C94C", "D64545"
M, I = "Montserrat", "Inter"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(int(1333.33 * U)), Emu(750 * U)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def new_slide(bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def text(slide, x, y, w, h, paras, align=None, shape=None):
    """paras: list of paragraphs; each is (text, font, size, color, bold) or a list of those runs."""
    if shape is None:
        shape = slide.shapes.add_textbox(Emu(int(x * U)), Emu(int(y * U)), Emu(int(w * U)), Emu(int(h * U)))
    tf = shape.text_frame
    tf.word_wrap = True
    for idx, p in enumerate(paras):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if align:
            para.alignment = align
        runs = p if isinstance(p, list) else [p]
        for t, font, size, color, bold in runs:
            r = para.add_run()
            r.text = t
            r.font.name = font
            r.font.size = Pt(round(size * FS, 1))
            r.font.color.rgb = rgb(color)
            r.font.bold = bold
        if not runs or runs[0][0] == "":
            pass
    return shape


def blank(size, color=SLATE):
    return ("", I, size, color, False)


def rect(slide, x, y, w, h, fill, rounded=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Emu(int(x * U)), Emu(int(y * U)), Emu(int(w * U)), Emu(int(h * U)))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(2 * FS)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if rounded:
        sh.adjustments[0] = 0.06
    return sh


def logo(s, color):
    text(s, 80, 14, 200, 26, [("AQUA SIGHT", M, 14, color, True)])


def footer(s, n, color):
    text(s, 1250, 705, 50, 32, [(str(n), I, 14, color, False)])


def title(s, t, color):
    text(s, 80, 40, 800, 60, [(t, M, 36, color, True)])


def body(lines, size, color=SLATE, font=I):
    return [(l, font, size, color, False) for l in lines]


# ---------- SLIDE 1: TITLE ----------
s = new_slide(NAVY)
text(s, 80, 220, 800, 300, [("AQUA SIGHT", M, 28, AQUA, True), blank(28), ("Eyes under the water.", M, 48, WHITE, True)])
text(s, 80, 520, 1000, 100, body(["Portable AI drowning detection for community pools",
                                  "[School Name]  |  [Presenter Names]"], 22, AQUA))
footer(s, 1, AQUA)

# ---------- SLIDE 2: SALES PITCH OUTLINE ----------
s = new_slide(BG)
title(s, "Sales Pitch Outline", NAVY)
logo(s, NAVY)
cards = [
    ("01", "The Problem", "About 11 Americans drown every day. It's the #1 cause of death for kids ages 1 to 4."),
    ("02", "The Market Gap", "About 1 in 3 U.S. public pools has been hit by the lifeguard shortage, and detection systems cost $50,000+."),
    ("03", "Our Solution", "Smart wristbands plus AI underwater cameras that warn lifeguards before a swimmer goes under."),
    ("04", "How It Works", "Three layers of detection: surface behavior, wristband sensors, and camera confirmation."),
    ("05", "Why We Win", "The only portable system combining wristbands and cameras. No construction. No video recorded."),
    ("06", "The Ask", "$150,000 for 15% equity ($1M post-money valuation)."),
]
for i, (num, head, desc) in enumerate(cards):
    x = 80 + (i % 3) * (364 + 24)
    y = 120 + (i // 3) * (250 + 20)
    rect(s, x, y, 364, 250, WHITE)
    text(s, x + 20, y + 20, 324, 210, [(num, M, 16, AQUA, True), blank(16),
                                       (head, M, 22, NAVY, True), (desc, I, 18, SLATE, False)])
text(s, 80, 665, 1170, 40, [("Market research: [ADD SURVEY RESULT, e.g., 'X of Y pool staff surveyed said they "
                             "would want this at their pool']", I, 14, SLATE, False)])
footer(s, 2, SLATE)

# ---------- SLIDE 3: PRODUCT OVERVIEW ----------
s = new_slide(BG)
logo(s, NAVY)
text(s, 80, 40, 1170, 80, [("Aqua Sight", M, 40, NAVY, True),
                           ("Three devices. Three layers of detection. One alarm that comes in time.", I, 18, SLATE, False)])
rect(s, 80, 130, 520, 420, WHITE)
text(s, 100, 150, 480, 380, body(["[ RENDER D: HERO SHOT ]", "", "1. Smart Wristband", "2. Underwater AI Camera",
                                  "3. Deck Hub", "", "(Clamps onto any pool in 15 mins. No construction.)"], 20, NAVY, M))
layers = [
    ("Layer 1: Surface Behavior AI", "Cameras recognize Instinctive Drowning Response: body vertical, head back, "
                                     "arms pressing down. Warns lifeguards BEFORE submersion."),
    ("Layer 2: Smart Wristband Sensors", "Tracks time underwater, sinking speed, and arm stoppage. "
                                         "Beginner & toddler bands trigger faster."),
    ("Layer 3: Multi-Camera Confirmation", "Cross-checks swimmer position & confirms head depth. "
                                           "Catches swimmers on bottom, band or no band."),
]
paras = [("How Aqua Sight detects drowning", M, 24, NAVY, True)]
for h, d in layers:
    paras += [blank(18), (h, I, 18, NAVY, True), (d, I, 18, SLATE, False)]
text(s, 630, 130, 620, 420, paras)
tiers = [("SAFE: Normal swimming", GREEN, WHITE),
         ("WATCH: Smartwatch buzzes 'Check lane 3'", YELLOW, NAVY),
         ("ALARM: Siren, strobe, exact spot on dashboard", RED, WHITE)]
tw = 1170 / 3
for k, (label, fill, tc) in enumerate(tiers):
    seg = rect(s, 80 + k * tw, 580, tw, 56, fill, rounded=False)
    seg.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(s, 0, 0, 0, 0, [(label, I, 15, tc, True)], align=PP_ALIGN.CENTER, shape=seg)
footer(s, 3, SLATE)

# ---------- SLIDE 4: PRODUCT COST ----------
s = new_slide(BG)
logo(s, NAVY)
title(s, "Product Cost", NAVY)
pricing = [("BUY", "$9,999", "one time + $99/mo", False),
           ("SUBSCRIBE", "$399", "per month, $0 upfront", True),
           ("RENT", "$299", "per weekend (meets/camps)", False)]
for j, (kind, price, sub, feat) in enumerate(pricing):
    px = 80 + j * (364 + 24)
    rect(s, px, 110, 364, 130, NAVY if feat else WHITE)
    text(s, px + 20, 120, 324, 110, [(kind, I, 14, AQUA if feat else SLATE, True),
                                     (price, M, 36, WHITE if feat else NAVY, True),
                                     (sub, I, 15, WHITE if feat else SLATE, False)])
text(s, 80, 260, 600, 270, [("Unit Cost to Make (One Full Kit):", I, 17, NAVY, True)] + body([
    "• 4 underwater AI camera pods ($395 ea): $1,580",
    "• 1 deck hub with siren & alarm: $200",
    "• 100 smart wristbands ($22 ea): $2,200",
    "• 2 band charging racks: $180",
    "• Carry case and packaging: $120"], 17) +
    [("TOTAL COST PER KIT: $4,280", I, 17, NAVY, True), ("(Plus $350 setup & staff training per pool)", I, 17, SLATE, False)])
text(s, 700, 260, 550, 270, [("Subscription Profitability:", I, 17, NAVY, True)] + body([
    "• $399 monthly price minus $40 cloud/warranty costs",
    "• Net: $359 per month per kit",
    "• Kit + setup ($4,630) paid back in ~13 months, then pure margin",
    "",
    "Purchase Margin: $5,369 per system (54%, after setup)",
    "Rental Margin: $224 profit per weekend"], 17))
text(s, 80, 540, 1170, 150, [("Competitive Benchmarks:", I, 15, NAVY, True)] + body([
    "• German municipal pool system: €50,000 for AI cameras",
    "• U.S. school district system: $253,837 total cost",
    "• Competitor WAVE: up to $399/mo (wearables only)",
    "• Aqua Sight: $399/mo (cameras + 100 bands included)",
    "Sources: KölnBäder; Fort Wayne Journal Gazette; WAVE published pricing."], 15))
footer(s, 4, SLATE)

# ---------- SLIDE 5: VIDEO ----------
s = new_slide(NAVY)
logo(s, AQUA)
title(s, "Our Pitch Video", WHITE)
rect(s, 200, 140, 933, 500, SLATE, line=AQUA)
text(s, 200, 340, 933, 100, body(["[ INSERT 3-MINUTE PITCH VIDEO HERE ]", "Featuring both team members"], 24, WHITE, M),
     align=PP_ALIGN.CENTER)
footer(s, 5, WHITE)

# ---------- SLIDE 6: INVESTOR REQUEST ----------
s = new_slide(BG)
logo(s, NAVY)
title(s, "Investor Request", NAVY)
text(s, 80, 120, 540, 540, [("$150,000", M, 60, TEAL, True), ("for 15% equity", M, 24, NAVY, True),
                            ("$1,000,000 post-money valuation", I, 20, SLATE, False), blank(17),
                            ("Milestones:", I, 17, NAVY, True)] + body([
    "• 5 pilot pools at camps and YMCAs, free through summer (April 2027), then converting to paid",
    "• 31 paying customers by end of 2027 (includes the 5 converted pilots)",
    "• Safety certification testing and growth round (2028)"], 17))
text(s, 650, 120, 600, 500, [("Where the money goes (Year 1) — Total: $250,375", M, 22, NAVY, True), blank(17)] + body([
    "• Hardware fleet (31 kits): $132,680 (53%)",
    "• Team: $36,000 (14%)",
    "• Engineering and testing: $29,000 (12%)",
    "• Marketing and trade shows: $18,100 (7%)",
    "• Insurance, legal, patents, software: $17,900 (7%)",
    "• Loan payments: $16,695 (7%)",
    "",
    "Funded by: $150k investment + $75k equipment loan + $5k founders + $20,375 customer revenue"], 17))
footer(s, 6, SLATE)

prs.save(sys.argv[1] if len(sys.argv) > 1 else "AQUA_SIGHT_Pitch_Deck.pptx")
